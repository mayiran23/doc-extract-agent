"""Test agent/tools/excel_writer.py — Excel 输出格式验证。"""

import os

import openpyxl
import pytest

from agent.tools.excel_writer import ExcelWriter


# ═══ 测试数据 ═══

SAMPLE_COLUMNS = [
    {"col_name": "铅含量", "semantic_type": "measured_value", "unit": "%"},
    {"col_name": "镉含量", "semantic_type": "measured_value", "unit": "%"},
    {"col_name": "pH", "semantic_type": "measured_value", "unit": ""},
    {"col_name": "温度(°C)", "semantic_type": "measured_value", "unit": "°C"},
]

SAMPLE_ROWS = [
    {
        "paper_id": "paper_001",
        "sequence_id": "A1",
        "values": {"铅含量": "8.5", "镉含量": "82.3", "pH": "7.2", "温度(°C)": "37"},
        "source_table": "Table 1",
        "source_figure": "",
        "citation": "Author. Title[J]. Journal, 2024, 1(1): 1-10. DOI: 10.1234/test",
        "notes": "测试备注：从Figure 3目测读取",
    },
    {
        "paper_id": "paper_001",
        "sequence_id": "A2",
        "values": {"铅含量": "8.5", "镉含量": "82.3", "pH": "7.2", "温度(°C)": "37"},
        "source_table": "Table 2",
        "source_figure": "",
        "citation": "Author. Title[J]. Journal, 2024, 1(1): 1-10. DOI: 10.1234/test",
        "notes": "",
    },
]


@pytest.fixture
def writer():
    """ExcelWriter 实例。"""
    return ExcelWriter()


def _get_cell(ws, row, col):
    """Helper — 获取单元格 (1-indexed)。"""
    return ws.cell(row=row, column=col)


class TestHeaderFormatting:
    """表头格式测试。"""

    def test_header_font_bold(self, writer, temp_dir):
        """表头应为白色粗体 Arial 9#。"""
        path = os.path.join(temp_dir, "test_header_font.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        cell = _get_cell(ws, 1, 1)
        assert cell.font.name == "Arial"
        assert cell.font.size == 9
        assert cell.font.bold is True
        assert cell.font.color.rgb == "00FFFFFF"

    def test_header_background_dark_blue(self, writer, temp_dir):
        """表头背景应为深蓝色 #2B6CB0。"""
        path = os.path.join(temp_dir, "test_header_bg.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        cell = _get_cell(ws, 1, 1)
        assert cell.fill.start_color.rgb == "002B6CB0"

    def test_header_has_thin_border(self, writer, temp_dir):
        """表头应有细线边框。"""
        path = os.path.join(temp_dir, "test_header_border.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        cell = _get_cell(ws, 1, 1)
        assert cell.border.left.style == "thin"
        assert cell.border.right.style == "thin"

    def test_header_center_aligned(self, writer, temp_dir):
        """表头应居中对齐。"""
        path = os.path.join(temp_dir, "test_header_align.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        cell = _get_cell(ws, 1, 1)
        assert cell.alignment.horizontal == "center"


class TestDataFormatting:
    """数据单元格格式测试。"""

    def test_data_font_arial_9(self, writer, temp_dir):
        """数据单元格应为 Arial 9#。"""
        path = os.path.join(temp_dir, "test_data_font.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        cell = _get_cell(ws, 2, 1)
        assert cell.font.name == "Arial"
        assert cell.font.size == 9

    def test_data_has_thin_border(self, writer, temp_dir):
        """数据单元格应有细线边框。"""
        path = os.path.join(temp_dir, "test_data_border.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        cell = _get_cell(ws, 2, 1)
        assert cell.border.left.style == "thin"

    def test_long_text_wrap_alignment(self, writer, temp_dir):
        """备注列应左对齐 + 自动换行。"""
        path = os.path.join(temp_dir, "test_long_text.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        notes_col = len(SAMPLE_COLUMNS) + 5  # 序号 + N列 + 数据来源 + 图表 + 图表2 + 备注
        cell = _get_cell(ws, 2, notes_col)
        assert cell.alignment.horizontal == "left"
        assert cell.alignment.wrap_text is True


class TestFreezePane:
    """冻结窗格测试。"""

    def test_first_row_frozen(self, writer, temp_dir):
        """首行应冻结。"""
        path = os.path.join(temp_dir, "test_freeze.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        assert ws.freeze_panes == "A2"


class TestColumnWidths:
    """列宽测试。"""

    def test_columns_have_width(self, writer, temp_dir):
        """所有列应有合适的宽度。"""
        path = os.path.join(temp_dir, "test_colwidth.xlsx")
        writer.write(SAMPLE_ROWS, SAMPLE_COLUMNS, path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active

        assert ws.column_dimensions["A"].width is not None
        assert ws.column_dimensions["A"].width >= 8


class TestEmptyOutput:
    """空数据输出测试。"""

    def test_empty_rows(self, writer, temp_dir):
        """零行数据应生成仅有表头的有效 Excel。"""
        path = os.path.join(temp_dir, "test_empty.xlsx")
        writer.write([], SAMPLE_COLUMNS, path)

        assert os.path.exists(path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        assert _get_cell(ws, 1, 1).value == "序号"
        assert _get_cell(ws, 2, 1).value is None
