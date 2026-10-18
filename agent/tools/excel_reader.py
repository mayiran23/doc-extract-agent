"""Excel 模板读取 — openpyxl 封装。

提供模板分析所需的：
- extract_headers: 读取第一行作为列名列表
- extract_sample_rows: 读取前 N 行数据供 LLM 理解列含义
"""

from typing import Dict, List

import openpyxl


def extract_headers(excel_path: str) -> List[str]:
    """从 Excel 文件第一行提取列名。

    Args:
        excel_path: Excel 文件的绝对路径。

    Returns:
        列名字符串列表（按从左到右的顺序）。

    Raises:
        ValueError: 文件格式不支持（如旧 .xls 格式需另存为 .xlsx）。
    """
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True)
    except Exception as e:
        raise ValueError(
            f"无法打开 Excel 文件。请确认：\n"
            f"  1. 文件是 .xlsx 格式（非旧版 .xls）\n"
            f"  2. 文件未被其他程序占用\n"
            f"  3. 文件没有损坏\n"
            f"原始错误: {e}"
        ) from e
    ws = wb.active

    headers = []
    for col in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=1, column=col).value
        headers.append(str(cell_value) if cell_value is not None else f"Column_{col}")

    wb.close()
    return headers


def extract_sample_rows(excel_path: str, n: int = 3) -> List[Dict[str, str]]:
    """从 Excel 文件中读取前 N 行数据，供 LLM 理解列的含义和上下文。

    Args:
        excel_path: Excel 文件的绝对路径。
        n: 读取的数据行数。

    Returns:
        [{列名: 值}, ...] 格式的样本数据。
    """
    headers = extract_headers(excel_path)

    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True)
    except Exception as e:
        raise ValueError(f"无法打开 Excel 文件: {e}") from e
    ws = wb.active

    sample_rows = []
    for row_idx in range(2, min(2 + n, ws.max_row + 1)):
        row_data = {}
        for col_idx, header in enumerate(headers, start=1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            row_data[header] = str(cell_value) if cell_value is not None else ""
        sample_rows.append(row_data)

    wb.close()
    return sample_rows
