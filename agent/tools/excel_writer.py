"""Excel 输出 — 含所有格式规则（继承自 SKILL.md）。

格式规范：
- 字体 Arial 9#
- 表头白色粗体，深蓝背景 #2B6CB0
- 细线边框，居中对齐
- 长文本列左对齐 + 自动换行
- 冻结首行
"""

from typing import Dict, List

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


class ExcelWriter:
    """将提取数据写入格式化 Excel 文件。"""

    # ── 格式常量 ──
    FONT_NAME = "Arial"
    FONT_SIZE = 9
    HEADER_BG_COLOR = "2B6CB0"
    HEADER_FONT_COLOR = "FFFFFF"

    # 长文本列 — 需要左对齐 + 自动换行
    LONG_TEXT_COLUMNS = {"数据来源", "备注", "图表", "图表2", "影响因素相关性分析"}

    def __init__(self):
        self._header_font = Font(
            name=self.FONT_NAME, size=self.FONT_SIZE, bold=True, color=self.HEADER_FONT_COLOR
        )
        self._header_fill = PatternFill(start_color=self.HEADER_BG_COLOR,
                                         end_color=self.HEADER_BG_COLOR, fill_type="solid")
        self._header_alignment = Alignment(horizontal="center", vertical="center")
        self._data_font = Font(name=self.FONT_NAME, size=self.FONT_SIZE)
        self._center_alignment = Alignment(horizontal="center", vertical="center")
        self._wrap_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        self._thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    def write(
        self,
        rows: List[Dict],
        columns: List[Dict],
        output_path: str,
        sheet_name: str = "检测报告数据提取",
    ) -> str:
        """将提取的数据写入格式化 Excel 文件。

        Args:
            rows: DataRow 列表，每行含 paper_id、sequence_id、values、notes 等。
            columns: ColumnDef 列表，定义列的顺序和元数据。
            output_path: 输出文件路径。
            sheet_name: 工作表名称。

        Returns:
            output_path 原样返回。
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 构建列顺序：序号 + 所有模板列 + 来源表格 + 来源位置 + 备注
        col_names = [c["col_name"] for c in columns]
        ordered_headers = ["序号"] + col_names + ["来源表格", "来源位置", "备注"]

        # ── 写入表头 ──
        for col_idx, header in enumerate(ordered_headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = self._header_font
            cell.fill = self._header_fill
            cell.alignment = self._header_alignment
            cell.border = self._thin_border

        # ── 写入数据行 ──
        for row_idx, row in enumerate(rows, start=2):
            # 序号列
            seq_cell = ws.cell(row=row_idx, column=1, value=row.get("sequence_id", ""))
            self._apply_data_style(seq_cell, is_long_text=False)

            # 模板列
            values = row.get("values", {})
            for col_offset, col_name in enumerate(col_names, start=1):
                cell = ws.cell(row=row_idx, column=col_offset + 1,
                               value=values.get(col_name, ""))
                is_long = col_name in self.LONG_TEXT_COLUMNS
                self._apply_data_style(cell, is_long_text=is_long)

            # 来源表格列
            table_col = len(col_names) + 2
            table_cell = ws.cell(row=row_idx, column=table_col,
                                 value=row.get("source_table", ""))
            self._apply_data_style(table_cell, is_long_text=False)

            # 来源位置列
            where_col = len(col_names) + 3
            where_cell = ws.cell(row=row_idx, column=where_col,
                                 value=row.get("source_figure", ""))
            self._apply_data_style(where_cell, is_long_text=False)

            # 备注列
            notes_col = len(col_names) + 4
            notes_cell = ws.cell(row=row_idx, column=notes_col, value=row.get("notes", ""))
            self._apply_data_style(notes_cell, is_long_text=True)

        # ── 调整列宽 ──
        self._auto_fit_columns(ws, ordered_headers)

        # ── 冻结首行 ──
        ws.freeze_panes = "A2"

        wb.save(output_path)
        return output_path

    def _apply_data_style(self, cell, is_long_text: bool = False):
        """为数据单元格应用统一样式。"""
        cell.font = self._data_font
        cell.border = self._thin_border
        if is_long_text:
            cell.alignment = self._wrap_alignment
        else:
            cell.alignment = self._center_alignment

    def _auto_fit_columns(self, ws, headers: List[str]):
        """根据表头长度自动调整列宽。"""
        for col_idx, header in enumerate(headers, start=1):
            char_width = sum(2 if ord(c) > 127 else 1 for c in header)
            width = max(8, min(char_width + 4, 40))
            ws.column_dimensions[get_column_letter(col_idx)].width = width
